"""Build Online Resource 2 (an Excel workbook) from the survey records.

Usage: python survey/scripts/build_online_resource2.py <output.xlsx>

Every data sheet is a verbatim copy of one CSV file in survey/records/. Integer and decimal
strings are written as numbers; all other values are written as text.
"""
import csv
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

R = Path(__file__).resolve().parents[1] / "records"
csv.field_size_limit(10_000_000)

TITLE = ("Trustworthy Artificial Intelligence in Metal Additive Manufacturing through a Critical Survey "
         "and a Controlled Benchmark")
IDENTIFICATION = [
    ("Online Resource 2", ""),
    ("Article", TITLE),
    ("Journal", "Journal of Intelligent Manufacturing"),
    ("Authors", "Ali Al Bataineh and Bandi Vamsi"),
    ("Corresponding author", "Ali Al Bataineh, Computer Information Science, Higher Colleges of Technology, "
                             "Abu Dhabi, United Arab Emirates; aalbataineh@hct.ac.ae"),
    ("", ""),
]

SHEETS = [
    ("queries", "query_log.csv",
     "Queries submitted to the Semantic Scholar bulk search interface on 15 October 2026, with request time, "
     "HTTP status and number of records retrieved"),
    ("records", "records_dedup.csv",
     "The 1,650 unique records after deduplication (identifier, DOI, title, authors, year, venue, publication "
     "type, citation count, open-access link, retrieving queries); abstracts are not redistributed"),
    ("screening_ledger", "screening_ledger.csv",
     "Screening decision, rule code and justification for every unique record"),
    ("screening_overrides", "screening_overrides.csv", "Screening decisions set by a recorded override"),
    ("spotcheck", "spotcheck.csv",
     "Check by the authors of a random sample of 94 screening decisions against the rules"),
    ("reviews", "reviews.csv", "Review and survey articles on AI or ML for metal AM published 2019 to 2026"),
    ("sample", "extraction_sample_order.csv",
     "Eligible studies with stratum, random position and sample flag (seed 20261015)"),
    ("evidence_extraction", "evidence_extraction.csv",
     "Evidence extracted from each sampled study, with access level and source locations"),
    ("eligibility_overrides", "eligibility_overrides.csv", "Sampled studies judged ineligible on reading"),
    ("verification", "verification.csv",
     "Second check, against the source, of a fixed set of coded fields (named in the field column) of every "
     "eligible study coded from full text and of two sampled studies later found ineligible, made by the "
     "authors"),
    ("coding_corrections", "coding_corrections.csv",
     "Corrections applied to the coded indicators, with reasons and source locations; a reason beginning "
     "Verification comes from the second check, Rule from the coding rules of Online Resource 1, Section S3, "
     "applied to all full-text studies, and Follows from marks an indicator that changes with a "
     "corrected one"),
    ("evidence_coded", "evidence_coded.csv", "Extracted evidence with the coded indicators used in the evidence map"),
    ("evidence_map", "evidence_map.csv", "Indicator counts per stratum and stratum-weighted proportions"),
    ("landscape", "landscape.csv", "Counts of processes, tasks, inputs and models over all eligible sampled studies"),
]

DICTIONARY = [
    ("Screening codes", "I1", "Eligible primary study: an implemented ML model applied to a metal AM process"),
    ("Screening codes", "E1", "Not metal additive manufacturing"),
    ("Screening codes", "E2", "No implemented ML model"),
    ("Screening codes", "E3", "Review, survey, editorial or perspective (routed to the review list)"),
    ("Screening codes", "E4", "Not a peer-reviewed journal article or full conference paper"),
    ("Screening codes", "E7", "Simulation without ML"),
    ("Screening codes", "E8", "Insufficient information in title and abstract for a decision"),
    ("Sample", "study_id", "Stratum letter (F, retrieved by a dimension query; G, general queries only) and "
                           "position in the random order of the stratum; a suffix _repNN marks the study taken "
                           "from reserve position NN to replace the study with that identifier"),
    ("Sample", "in_sample", "True for the first 60 positions of the random order of each stratum (the drawn "
                            "sample); a drawn study judged ineligible on reading keeps True and is replaced by the "
                            "next unused reserve position (61 onward), whose row keeps False; in the evidence sheets "
                            "a reserve study carries the identifier of the study it replaced with the suffix _repNN "
                            "(for example F061 appears as F010_rep61)"),
    ("Sample", "eligible", "Value recorded at extraction: Y, or N:Ek with Ek the screening code that applies. The "
                           "final eligibility of a sampled study is given by eligibility_overrides.csv, which "
                           "excludes six studies, three of which were recorded Y at extraction (F003, F010, G024)"),
    ("Sample", "access_level", "Version read: full text (open access or publisher page) or abstract only; "
                               "studies coded from the abstract only are excluded from the evidence map"),
    ("Spot check", "rule_code_at_check, screener_code", "Code given by the screening rules when the sample was "
                                                        "checked, and code assigned by the authors"),
    ("Spot check", "agree", "yes when the two codes are equal, no otherwise"),
    ("Spot check", "action", "Override or rule change made after a disagreement; empty when none was needed"),
    ("Evidence map", "denominator", "For every indicator, the number of studies coded from full text in the "
                                    "stratum (F_n, G_n)"),
    ("Evidence map", "True / False", "True when the study reported evidence of the practice; False when it did "
                                     "not report it, which does not establish that the practice was absent"),
    ("Evidence map", "weighted_prop", "Stratum-weighted proportion, with weights equal to the stratum sizes "
                                      "(554 and 652)"),
    ("Indicator", "split_random_only", "Random, shuffled or stratified assignment of samples to the subsets or folds "
                                       "stated, and neither a grouped partition nor a held-out-condition test"),
    ("Indicator", "split_grouped", "Samples grouped by part, specimen or layer block"),
    ("Indicator", "split_heldout_condition", "Test data from process parameters, geometries, builds, machines, "
                                             "materials, external datasets or experiments withheld from training "
                                             "by design so as to differ from the training conditions; test samples "
                                             "drawn from the same design or sampling as the training samples are "
                                             "classified by how they were assigned, even when each has its own "
                                             "parameter combination"),
    ("Indicator", "split_beyond_random", "Grouped partition or test on held-out conditions"),
    ("Indicator", "split_not_reported", "No partition reported, or a partition given only by proportions, counts or "
                                        "a number of folds without stating how samples were assigned"),
    ("Indicator", "heldout_machine, heldout_material, heldout_build_or_experiment",
     "Test data from a machine, material, or build or experiment not used for training"),
    ("Indicator", "simulation_data_only", "All data from simulation"),
    ("Indicator", "leakage_high, leakage_unclear, leakage_low",
     "Rated risk that correlated samples appear in both training and test data, for the first reported "
     "evaluation of the model on data not used for training"),
    ("Indicator", "calibration_any", "A calibration measure or calibration method reported"),
    ("Indicator", "uncertainty_any", "Uncertainty quantified"),
    ("Indicator", "uncertainty_validated", "Uncertainty compared with errors, coverage or independent measurements"),
    ("Indicator", "robustness_any, robustness_shift, robustness_perturbation",
     "Any robustness test; test under changed conditions; defined input perturbation at test time "
     "(training augmentation not counted)"),
    ("Indicator", "explanation_any", "An explanation method used"),
    ("Indicator", "explanation_validated", "Explanation compared with annotations, measured properties or another "
                                           "independent reference"),
    ("Indicator", "explanation_qualitative_only", "Explanations inspected qualitatively only"),
    ("Indicator", "data_public, code_public", "Data public fully or in part; code public"),
    ("Indicator", "repeats_reported", "Repeated training runs, seeds or cross-validation folds whose results are "
                                      "reported; cross-validation used only to select hyperparameters, without "
                                      "reported fold results, is not counted"),
    ("Indicator", "variability_reported", "Standard deviation, standard error, range, interval, box plot or table of "
                                          "a performance measure across repeated training runs, seeds or "
                                          "cross-validation folds; spreads of measured data, predictive uncertainty "
                                          "intervals, tests comparing means, spreads of errors across the samples "
                                          "of one test set and spreads across configurations each trained once are "
                                          "not counted"),
    ("Indicator", "compute_any, compute_inference", "Any computational evidence; inference time or throughput"),
    ("Indicator", "realtime_claim, realtime_in_loop, closed_loop",
     "Real-time suitability claimed; timing reported for a model running during an actual or physically "
     "emulated process; closed-loop control"),
    ("Verification", "field", "Coded field checked. Field names follow two styles, with spaces or with "
                              "underscores, for the same fields (for example closed loop and closed_loop; split "
                              "category and split; ELIGIBILITY and eligibility); fields marked (derived), input_cat "
                              "and task_cat are landscape categories"),
    ("Verification", "field and indicator",
     "explanation used: explanation_any; explanation checked: explanation_validated; uncertainty quantified: "
     "uncertainty_any; uncertainty validated: uncertainty_validated; inference timing: compute_inference; timing in "
     "process: realtime_in_loop; real-time claim: realtime_claim; closed loop: closed_loop; split category: the split "
     "indicators; leakage_risk: the leakage indicators; robustness: the robustness indicators; calibration: "
     "calibration_any; repeats: repeats_reported; variability: variability_reported; held-out machine: "
     "heldout_machine; held-out material: heldout_material; held-out build or experiment: "
     "heldout_build_or_experiment; data public: data_public; code public: code_public; "
     "simulation-only data: simulation_data_only"),
    ("Verification", "original_value, checked_value",
     "Value of the first extraction and value found by the authors in the source at the second check; same "
     "when the source confirmed "
     "the first value as recorded; NA when the field does not apply to the study (for example uncertainty "
     "validated when no uncertainty was quantified). Y/N and True/False are equivalent value conventions"),
    ("Verification", "agree", "Y when the checked value agrees in substance with the first value, N when it does "
                              "not; NA in one eligibility row (F031), where the check referred the study to rule E7 "
                              "instead of recording agreement. F031 implements a physics-informed neural network on "
                              "simulated data and was kept as eligible, since E7 applies only to simulation "
                              "without ML"),
    ("Verification", "locator, access", "Location of the supporting passage and the version of the study read"),
    ("Coding corrections", "reason", "Verification: value found in the second check against the source; Rule: "
                                     "value set by the coding rules of Online Resource 1, Section S3, applied to all "
                                     "full-text studies, which can differ from the value recorded in the second "
                                     "check; Follows from: indicator that changes with a corrected indicator"),
    ("Landscape", "process_cat, task_cat, input_cat, model_cat",
     "Categories of process, task, input and model over all eligible sampled studies"),
]

INT = re.compile(r"-?(0|[1-9][0-9]*)")
DEC = re.compile(r"-?[0-9]+\.[0-9]+([eE]-?[0-9]+)?")


def cell(value):
    if INT.fullmatch(value):
        return int(value)
    if DEC.fullmatch(value):
        return float(value)
    return value


def main():
    out = Path(sys.argv[1])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    for a, b in IDENTIFICATION:
        ws.append([a, b])
    ws.append(["Sheet", "Content"])
    for name, _, text in SHEETS:
        ws.append([name, text])
    ws.append(["dictionary", "Screening codes, sample identifiers and the definition of every coded indicator"])
    bold = Font(bold=True)
    ws["A1"].font = bold
    ws.cell(row=len(IDENTIFICATION) + 1, column=1).font = bold
    ws.cell(row=len(IDENTIFICATION) + 1, column=2).font = bold
    for name, fname, _ in SHEETS:
        sh = wb.create_sheet(name)
        with open(R / fname, newline="", encoding="utf-8") as f:
            for i, row in enumerate(csv.reader(f)):
                sh.append(row if i == 0 else [cell(v) for v in row])
        for c in sh[1]:
            c.font = bold
    sh = wb.create_sheet("dictionary")
    sh.append(["Group", "Item", "Definition"])
    for row in DICTIONARY:
        sh.append(list(row))
    for c in sh[1]:
        c.font = bold
    stamp = dt.datetime(2026, 10, 15)
    wb.properties.creator = "Ali Al Bataineh"
    wb.properties.lastModifiedBy = "Ali Al Bataineh"
    wb.properties.title = "Online Resource 2"
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.save(out)
    print("wrote", out, "with", len(wb.sheetnames), "sheets")


if __name__ == "__main__":
    main()
